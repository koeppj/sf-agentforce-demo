from pathlib import Path
import ast,json,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
P=Path(__file__).resolve().parent
fields=json.loads((P/'field_map.json').read_text())
d=Document();s=d.sections[0]
s.page_width=Inches(8.5);s.page_height=Inches(11)
s.top_margin=Inches(.65);s.bottom_margin=Inches(.65);s.left_margin=s.right_margin=Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 st=d.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.space_after=Pt(6)
d.styles['Normal'].font.size=Pt(10)
d.styles['Normal'].paragraph_format.line_spacing=1.05
d.styles['Title'].font.size=Pt(22)
d.styles['Heading 1'].font.size=Pt(16)
d.styles['Heading 2'].font.size=Pt(12)
footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('Allegiance Mobile Health  |  ');r.font.size=Pt(8)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
def p(txt,size=None):
 a=d.add_paragraph(txt)
 if size:
  for r in a.runs:r.font.size=Pt(size)
 return a

def heading(txt):d.add_heading(txt,level=1)
def field(f):
 a=d.add_paragraph();a.paragraph_format.space_after=Pt(9);a.paragraph_format.keep_together=True
 r=a.add_run(f['label']+(' *' if f['required'] else ''));r.bold=True;r.font.size=Pt(9)
 a.add_run().add_break()
 r=a.add_run('{{'+f['json_path']+'}}');r.underline=True;r.font.size=Pt(11)
 # Complete each tag in one ordinary Word text run, with no content control or merge field.

p('ALLEGIANCE MOBILE HEALTH',10)
d.add_paragraph('New hire form',style='Title')
p('Complete the applicable sections to request a new hire or personnel change. The completed form begins the approval process. You will be notified if the request is declined.')
heading('New hire instructions')
source=ast.parse((P/'build_template.py').read_text())
steps=next(ast.literal_eval(n.value) for n in source.body if isinstance(n,ast.Assign) and any(isinstance(t,ast.Name) and t.id=='steps' for t in n.targets))
for n,txt in steps:
 a=d.add_paragraph();a.paragraph_format.space_after=Pt(7)
 a.add_run(n+'. ').bold=True;a.add_run(txt)
 for r in a.runs:r.font.size=Pt(9)

d.add_page_break();heading('Employee and assignment')
p('Fields marked with an asterisk (*) are required.',9)
for f in fields:
 if f['page']==2:field(f)
p('For region, district and station, enter NA if there are no changes. Notify HRHelp@allmh.com immediately if the start date, FT/PRN status or station changes, or if the employee will be a never worked.',9)
p('Contact HRHelp@allmh.com if a new station or job title needs to be added.',9)

d.add_page_break();heading('Role certification and pay')
for f in fields:
 if f['page']==3:field(f)
p('For the new manager, enter NA if there are no changes.',9)
p('If certified in another state or holding an expired certification predating the DSHS date, provide the prior certification details and send supporting documentation to HRHelp@allmh.com.',9)
p('If above minimum, a comment is required. If above maximum, approval documentation from Dan Gillespie is required and must be included with supporting documentation.',9)

d.add_page_break();heading('Notes and authorization')
p('Only authorized managers are allowed to submit forms. Submissions from unauthorized individuals may be rejected.',10)
field(next(f for f in fields if f['name']=='notes'))
p('Provide supporting documentation separately, including applicable employee communications, transfers, certifications or pay changes. If received later, email HRHelp@allmh.com.',9)
for f in fields:
 if f['name'].startswith('submitter.'):field(f)
p('Attestation: I confirm that the information provided is accurate to the best of my knowledge and that I am authorized to request or make this change.',9)
d.add_heading('Record references',level=2)
for f in fields:
 if f['name'].startswith(('case.','schema.')):field(f)
for root in [d._element, d.styles.element]:
 for border in list(root.iter(qn('w:pBdr'))):border.getparent().remove(border)
file=P/'new_hire_box_docgen.docx';d.save(file)
check=Document(file);tags=[r.text for a in check.paragraphs for r in a.runs if '{{' in r.text]
assert len(tags)==30
assert set(tags)=={'{{'+f['json_path']+'}}' for f in fields}
assert all(r.underline for a in check.paragraphs for r in a.runs if '{{' in r.text)
assert not check.tables
assert '{{user_data.' not in '\n'.join(tags)
print('PASS: 30 exact JSON-path tags, each in one underlined run; no tables or boxes.')
# Local sample substitution only for visual QA; not a Box Doc Gen execution.
data=json.loads((P/'user_data.example.json').read_text())
for a in check.paragraphs:
 for r in a.runs:
  if r.text in tags:
   val=data
   for k in r.text[2:-2].split('.'):val=val[k]
   r.text=str(val).lower() if isinstance(val,bool) else str(val)
qa=P/'word_qa';qa.mkdir(exist_ok=True)
check.save(qa/'sample.docx')
